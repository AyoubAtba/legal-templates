"""
Convert the Markdown legal documents in output/ to plain .docx files in output/docx/.

Plain formatting only: Word's built-in Heading 1-3 styles, bold/italic runs,
bulleted and numbered lists, simple tables, and paragraphs. No colors, fonts,
or custom styling.

Usage:
    python3 scripts/build_docx.py            # converts every output/*.md
    python3 scripts/build_docx.py FILE.md    # converts the given file(s)

Requires python-docx (see scripts/requirements.txt).
"""

import datetime
import io
import re
import sys
import zipfile
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "output" / "docx"

# Fixed timestamp for document metadata and zip entries, so that rebuilding
# unchanged Markdown yields identical bytes.
FIXED_TIME = datetime.datetime(2026, 1, 1)

INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def add_inline(paragraph, text):
    """Add text to a paragraph, honoring **bold**, *italic*, and `code` markers."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith("`") and part.endswith("`"):
            paragraph.add_run(part[1:-1])
        else:
            paragraph.add_run(part)


def flush_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-+:?", c or "-") for c in r)]
    if not cells:
        return
    ncols = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(ncols):
            cell = table.cell(i, j)
            cell.text = ""
            add_inline(cell.paragraphs[0], row[j] if j < len(row) else "")
            if i == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()


def convert(md_path: Path, docx_path: Path):
    doc = Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()
    table_rows = []
    para_buf = []

    def flush_para():
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para_buf))
            para_buf.clear()

    for line in lines:
        stripped = line.strip()

        if stripped.startswith("|"):
            flush_para()
            table_rows.append(stripped)
            continue
        if table_rows:
            flush_table(doc, table_rows)
            table_rows = []

        if not stripped:
            flush_para()
            continue
        if stripped == "---":
            flush_para()
            doc.add_paragraph("_" * 40)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_para()
            level = min(len(m.group(1)), 3)
            add_inline(doc.add_heading("", level=level), m.group(2))
            continue
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            flush_para()
            style = "List Bullet 2" if len(m.group(1)) >= 2 else "List Bullet"
            add_inline(doc.add_paragraph(style=style), m.group(2))
            continue
        m = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if m:
            flush_para()
            style = "List Number 2" if len(m.group(1)) >= 2 else "List Number"
            add_inline(doc.add_paragraph(style=style), m.group(2))
            continue
        para_buf.append(stripped)

    flush_para()
    if table_rows:
        flush_table(doc, table_rows)

    props = doc.core_properties
    props.author = "Chexbill LLC"
    props.last_modified_by = "Chexbill LLC"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.revision = 1
    props.title = md_path.stem

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    save_reproducibly(doc, docx_path)


def save_reproducibly(doc, docx_path):
    """Save doc with fixed zip entry timestamps and attributes.

    python-docx stamps each zip entry with the current time, so an unchanged
    document would otherwise produce different bytes on every build.
    """
    buf = io.BytesIO()
    doc.save(buf)
    with zipfile.ZipFile(buf) as src, zipfile.ZipFile(docx_path, "w") as dst:
        for entry in src.infolist():
            info = zipfile.ZipInfo(entry.filename, date_time=FIXED_TIME.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3  # Unix, regardless of the build platform
            info.external_attr = 0o644 << 16
            dst.writestr(info, src.read(entry))


def main(argv):
    sources = [Path(a).resolve() for a in argv] or sorted((ROOT / "output").glob("*.md"))
    for src in sources:
        dest = OUT_DIR / (src.stem + ".docx")
        convert(src, dest)
        print(f"{src.relative_to(ROOT)} -> {dest.relative_to(ROOT)}")


if __name__ == "__main__":
    main(sys.argv[1:])
